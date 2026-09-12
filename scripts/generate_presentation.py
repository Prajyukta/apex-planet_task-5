from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


SLIDES = [
    ("Apexplanet Retail Intelligence", "Final Reporting, Pipeline Deployment & Executive Presentation\nTask 5 Capstone"),
    ("The business problem", "Manual reporting delays action and fragments KPI definitions.\nDecision question: what changed, where is margin leaking, and which customers need intervention?"),
    ("Data and operating contract", "11 source fields across transaction, customer, product, geography, commercial, and churn domains.\nControls: schema validation, UTC dates, numeric bounds, deduplication."),
    ("EDA: performance signals", "Revenue, profit margin, AOV, active customers, and units sold.\nDiscoveries: mix, concentration, discount-margin relationship, repeat behavior, and risk concentration."),
    ("Executive dashboard", "Overview | Regional drill-through | Customer risk | Tooltip microchart\nFilters: date, region, category. Source: processed CSV artifacts."),
    ("Statistics and machine learning", "ANOVA, t-tests, and confidence intervals.\nRFM K-Means segments and explainable churn-risk score.\nARIMA 92% accuracy: validate on the approved holdout before publishing."),
    ("Recommendations", "Immediate: retain high-value/high-risk customers and investigate margin outliers.\nMid-term: category playbooks and experiments.\nLong-term: monitored propensity and forecast-led planning."),
    ("Challenges and controls", "Batch quality variation, sparse history, metric governance, and dashboard ownership.\nControls: data contract, secrets, logs, versioned artifacts, and scheduled refresh."),
    ("Future scope", "Move CSV handoff to governed warehouse tables.\nAdd freshness, drift, anomaly, and model monitoring.\nConnect causal uplift testing to CRM activation."),
    ("Thank you", "A reproducible daily analytics operating loop.\nRepository: <URL>   Dashboard: <URL>   Demo: <URL>"),
]


def build_deck(output: Path) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    for index, (title, body) in enumerate(SLIDES, start=1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = RGBColor(11, 39, 53)
        accent = slide.shapes.add_shape(1, Inches(0.7), Inches(0.75), Inches(0.18), Inches(5.9))
        accent.fill.solid()
        accent.fill.fore_color.rgb = RGBColor(242, 153, 74)
        accent.line.fill.background()
        title_box = slide.shapes.add_textbox(Inches(1.2), Inches(1.05), Inches(11.3), Inches(1.4))
        title_frame = title_box.text_frame
        title_frame.text = title
        title_frame.paragraphs[0].font.size = Pt(30)
        title_frame.paragraphs[0].font.bold = True
        title_frame.paragraphs[0].font.color.rgb = RGBColor(247, 247, 242)
        body_box = slide.shapes.add_textbox(Inches(1.25), Inches(2.65), Inches(10.8), Inches(2.5))
        body_frame = body_box.text_frame
        body_frame.text = body
        body_frame.word_wrap = True
        body_frame.paragraphs[0].font.size = Pt(21)
        body_frame.paragraphs[0].font.color.rgb = RGBColor(220, 232, 229)
        footer = slide.shapes.add_textbox(Inches(1.25), Inches(6.8), Inches(10), Inches(0.3))
        footer.text_frame.text = f"APEXPLANET  /  TASK 5                                      {index:02d}"
        footer.text_frame.paragraphs[0].font.size = Pt(9)
        footer.text_frame.paragraphs[0].font.color.rgb = RGBColor(161, 190, 186)
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)


if __name__ == "__main__":
    build_deck(Path(__file__).resolve().parents[1] / "presentation" / "presentation_slides.pptx")